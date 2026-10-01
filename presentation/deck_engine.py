"""Tiny slide engine: a slide is a list of positioned elements. The same elements are
written to .pptx (python-pptx) and drawn to a PNG preview (PIL) with real Calibri/Cambria
metrics, so text fit and overlaps can be checked without PowerPoint or LibreOffice.
"""
from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

W, H = 13.333, 7.5
INK, INK2, MUTED = "0B0B0B", "52514E", "898781"
ACCENT, DARK, TINT, LINE = "2A78D6", "13223A", "EEF4FB", "E1E0D9"
WHITE, ICE = "FFFFFF", "CADCFC"
FONTS = {"Calibri": "C:/Windows/Fonts/calibri.ttf", "Calibri-b": "C:/Windows/Fonts/calibrib.ttf",
         "Cambria": "C:/Windows/Fonts/cambria.ttc", "Cambria-b": "C:/Windows/Fonts/cambriab.ttf"}


@dataclass
class Text:
    x: float
    y: float
    w: float
    h: float
    paras: list            # list of str or (str, dict) ; dict keys: bullet, bold, size, color, italic
    size: float = 16
    color: str = INK
    font: str = "Calibri"
    bold: bool = False
    align: str = "left"
    anchor: str = "top"
    space_after: float = 6


@dataclass
class Img:
    x: float
    y: float
    w: float
    h: float
    path: str


@dataclass
class Box:
    x: float
    y: float
    w: float
    h: float
    fill: str = TINT
    rounded: bool = True


@dataclass
class Table:
    x: float
    y: float
    w: float
    rows: list
    col_w: list
    size: float = 11
    row_h: float = 0.3
    highlight_row: int | None = None


@dataclass
class Slide:
    elements: list
    notes: str = ""
    bg: str = WHITE
    title: str = ""


def _rgb(h):
    return RGBColor.from_string(h)


# ------------------------------------------------------------------------------------------ measuring
def _font(name, bold, size_pt, scale):
    key = f"{name}-b" if bold else name
    return ImageFont.truetype(FONTS.get(key, FONTS["Calibri"]), max(1, int(round(size_pt * scale))))


def wrap_lines(text, font, width_px):
    lines = []
    for raw in text.split("\n"):
        words, cur = raw.split(" "), ""
        for w_ in words:
            trial = (cur + " " + w_).strip()
            if font.getlength(trial) <= width_px or not cur:
                cur = trial
            else:
                lines.append(cur)
                cur = w_
        lines.append(cur)
    return lines


def text_height(t: Text, scale=1.0):
    """Height in points of the wrapped text (python-pptx default insets: 0.1in left/right, 0.05in top/bottom)."""
    inner_w = (t.w - 0.2) * 72 * scale
    total = 0.0
    for p in t.paras:
        s, o = (p, {}) if isinstance(p, str) else p
        size = o.get("size", t.size)
        f = _font(t.font, o.get("bold", t.bold), size, scale)
        indent = 0.25 * 72 * scale if o.get("bullet") else 0
        n = len(wrap_lines(s, f, inner_w - indent))
        total += n * size * 1.2 * scale + t.space_after * scale
    return total / scale + 7.2        # + top/bottom insets


def check(slides, deck):
    problems = []
    for i, s in enumerate(slides, 1):
        for e in s.elements:
            if isinstance(e, Text):
                need = text_height(e) / 72
                if need > e.h + 0.02:
                    problems.append(f"{deck} slide {i}: text needs {need:.2f}in > box {e.h:.2f}in: {str(e.paras[0])[:50]}")
            if hasattr(e, "x") and not isinstance(e, Table):
                if e.x < 0.3 or e.y < 0.2 or e.x + e.w > W - 0.3 + 1e-6 or e.y + e.h > H - 0.15 + 1e-6:
                    problems.append(f"{deck} slide {i}: element outside safe area {type(e).__name__} {e.x, e.y, e.w, e.h}")
    return problems


# ------------------------------------------------------------------------------------------ pptx
def _fit(path, x, y, w, h):
    iw, ih = Image.open(path).size
    r = min(w / iw, h / ih)
    fw, fh = iw * r, ih * r
    return x + (w - fw) / 2, y + (h - fh) / 2, fw, fh


def write_pptx(slides, out: Path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    blank = prs.slide_layouts[6]
    for s in slides:
        sl = prs.slides.add_slide(blank)
        bg = sl.background.fill
        bg.solid()
        bg.fore_color.rgb = _rgb(s.bg)
        for e in s.elements:
            if isinstance(e, Box):
                shp = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if e.rounded else MSO_SHAPE.RECTANGLE,
                                          Inches(e.x), Inches(e.y), Inches(e.w), Inches(e.h))
                shp.fill.solid()
                shp.fill.fore_color.rgb = _rgb(e.fill)
                shp.line.fill.background()
                if e.rounded:
                    shp.adjustments[0] = 0.08
            elif isinstance(e, Img):
                x, y, w, h = _fit(e.path, e.x, e.y, e.w, e.h)
                sl.shapes.add_picture(str(e.path), Inches(x), Inches(y), Inches(w), Inches(h))
            elif isinstance(e, Text):
                tb = sl.shapes.add_textbox(Inches(e.x), Inches(e.y), Inches(e.w), Inches(e.h))
                tf = tb.text_frame
                tf.word_wrap = True
                tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[e.anchor]
                for k, p in enumerate(e.paras):
                    s_, o = (p, {}) if isinstance(p, str) else p
                    para = tf.paragraphs[0] if k == 0 else tf.add_paragraph()
                    para.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER}[e.align]
                    para.space_after = Pt(e.space_after)
                    if o.get("bullet"):
                        _bullet(para)
                    run = para.add_run()
                    run.text = s_
                    run.font.name = e.font
                    run.font.size = Pt(o.get("size", e.size))
                    run.font.bold = o.get("bold", e.bold)
                    run.font.italic = o.get("italic", False)
                    run.font.color.rgb = _rgb(o.get("color", e.color))
            elif isinstance(e, Table):
                n, m = len(e.rows), len(e.rows[0])
                gt = sl.shapes.add_table(n, m, Inches(e.x), Inches(e.y), Inches(e.w), Inches(e.row_h * n)).table
                for j, cw in enumerate(e.col_w):
                    gt.columns[j].width = Inches(cw)
                for r_ in range(n):
                    gt.rows[r_].height = Inches(e.row_h)
                    for c_ in range(m):
                        cell = gt.cell(r_, c_)
                        cell.margin_left = cell.margin_right = Inches(0.06)
                        cell.margin_top = cell.margin_bottom = Inches(0.02)
                        cell.fill.solid()
                        head = r_ == 0
                        hl = e.highlight_row is not None and r_ == e.highlight_row
                        cell.fill.fore_color.rgb = _rgb(DARK if head else (TINT if hl else WHITE))
                        tf = cell.text_frame
                        tf.paragraphs[0].text = ""
                        run = tf.paragraphs[0].add_run()
                        run.text = str(e.rows[r_][c_])
                        run.font.size = Pt(e.size)
                        run.font.name = "Calibri"
                        run.font.bold = head or hl
                        run.font.color.rgb = _rgb(WHITE if head else INK)
                        tf.paragraphs[0].alignment = PP_ALIGN.LEFT if c_ == 0 else PP_ALIGN.CENTER
        if s.notes:
            sl.notes_slide.notes_text_frame.text = s.notes
    prs.save(out)


def _bullet(para):
    from pptx.oxml.ns import qn
    pPr = para._p.get_or_add_pPr()
    pPr.set("marL", str(Inches(0.25)))
    pPr.set("indent", str(-Inches(0.2)))
    for tag in ("a:buNone", "a:buChar", "a:buAutoNum"):
        for el in pPr.findall(qn(tag)):
            pPr.remove(el)
    bu = pPr.makeelement(qn("a:buChar"), {"char": "\u2022"})
    pPr.append(bu)


# ------------------------------------------------------------------------------------------ preview
def preview(slides, out_prefix: Path, dpi=80):
    paths = []
    for i, s in enumerate(slides, 1):
        im = Image.new("RGB", (int(W * dpi), int(H * dpi)), "#" + s.bg)
        d = ImageDraw.Draw(im)
        sc = dpi / 72
        for e in s.elements:
            if isinstance(e, Box):
                d.rounded_rectangle([e.x * dpi, e.y * dpi, (e.x + e.w) * dpi, (e.y + e.h) * dpi], radius=8, fill="#" + e.fill)
            elif isinstance(e, Img):
                x, y, w, h = _fit(e.path, e.x, e.y, e.w, e.h)
                pic = Image.open(e.path).convert("RGB").resize((max(1, int(w * dpi)), max(1, int(h * dpi))))
                im.paste(pic, (int(x * dpi), int(y * dpi)))
            elif isinstance(e, Table):
                for r_, row in enumerate(e.rows):
                    cx = e.x
                    for c_, val in enumerate(row):
                        box = [cx * dpi, (e.y + r_ * e.row_h) * dpi, (cx + e.col_w[c_]) * dpi, (e.y + (r_ + 1) * e.row_h) * dpi]
                        fill = "#" + (DARK if r_ == 0 else (TINT if r_ == e.highlight_row else WHITE))
                        d.rectangle(box, fill=fill, outline="#" + LINE)
                        f = _font("Calibri", r_ == 0 or r_ == e.highlight_row, e.size, sc)
                        d.text((box[0] + 4, box[1] + 3), str(val), font=f, fill="#" + (WHITE if r_ == 0 else INK))
                        cx += e.col_w[c_]
            elif isinstance(e, Text):
                y = (e.y + 0.05) * dpi
                inner = (e.w - 0.2) * 72 * sc
                blocks = []
                for p in e.paras:
                    s_, o = (p, {}) if isinstance(p, str) else p
                    size = o.get("size", e.size)
                    f = _font(e.font, o.get("bold", e.bold), size, sc)
                    ind = 0.25 * dpi if o.get("bullet") else 0
                    lines = wrap_lines(s_, f, inner - ind)
                    blocks.append((lines, f, size, ind, o))
                total = sum(len(b[0]) * b[2] * 1.2 * sc + e.space_after * sc for b in blocks)
                if e.anchor == "middle":
                    y = e.y * dpi + (e.h * dpi - total) / 2
                elif e.anchor == "bottom":
                    y = (e.y + e.h) * dpi - total
                for lines, f, size, ind, o in blocks:
                    for k, ln in enumerate(lines):
                        tx = (e.x + 0.1) * dpi + ind
                        if e.align == "center":
                            tx = (e.x + e.w / 2) * dpi - f.getlength(ln) / 2
                        if o.get("bullet") and k == 0:
                            d.text(((e.x + 0.1) * dpi + 0.05 * dpi, y), "\u2022", font=f, fill="#" + o.get("color", e.color))
                        d.text((tx, y), ln, font=f, fill="#" + o.get("color", e.color))
                        y += size * 1.2 * sc
                    y += e.space_after * sc
                d.rectangle([e.x * dpi, e.y * dpi, (e.x + e.w) * dpi, (e.y + e.h) * dpi], outline="#F0D0D0")
        p = Path(f"{out_prefix}-{i:02d}.png")
        im.save(p)
        paths.append(p)
    return paths


# ------------------------------------------------------------------------------------------ layouts
def title_slide(title, subtitle, footer, notes=""):
    return Slide(bg=DARK, notes=notes, title=title, elements=[
        Text(0.8, 2.0, 11.7, 1.9, [title], size=40, color=WHITE, font="Cambria", bold=True, anchor="bottom"),
        Text(0.8, 4.05, 11.7, 1.1, [subtitle], size=20, color=ICE),
        Text(0.8, 6.3, 11.7, 0.5, [footer], size=12, color=ICE)])


def header(title, kicker=None):
    els = [Text(0.6, 0.35, 12.1, 0.85, [title], size=30, color=INK, font="Cambria", bold=True, anchor="middle")]
    if kicker:
        els.append(Text(0.6, 1.15, 12.1, 0.5, [kicker], size=15, color=INK2))
    return els


def bullets(items, size=16):
    out = []
    for it in items:
        if isinstance(it, tuple):
            out.append((it[0], {"bullet": True, **it[1]}))
        else:
            out.append((it, {"bullet": True}))
    return out


def text_image(title, kicker, items, img, notes="", split=5.3, size=16):
    els = header(title, kicker)
    top = 1.75 if kicker else 1.4
    els.append(Text(0.6, top, split, H - top - 0.45, bullets(items, size), size=size, space_after=8))
    els.append(Img(0.6 + split + 0.3, top, W - 0.6 - (0.6 + split + 0.3), H - top - 0.45, str(img)))
    return Slide(elements=els, notes=notes, title=title)


def stats_row(stats, y=1.85, h=1.45, x0=0.6, width=None):
    width = width or (W - 1.2)
    n = len(stats)
    gap = 0.3
    w = (width - gap * (n - 1)) / n
    els = []
    for i, (big, label) in enumerate(stats):
        x = x0 + i * (w + gap)
        els.append(Box(x, y, w, h))
        els.append(Text(x + 0.1, y + 0.08, w - 0.2, 0.75, [big], size=30, color=ACCENT, bold=True, anchor="bottom"))
        els.append(Text(x + 0.1, y + 0.83, w - 0.2, h - 0.88, [label], size=12, color=INK2))
    return els


def stats_image(title, kicker, stats, img, caption=None, notes=""):
    els = header(title, kicker) + stats_row(stats)
    top = 3.55
    bottom = H - (0.8 if caption else 0.45)
    els.append(Img(0.6, top, W - 1.2, bottom - top, str(img)))
    if caption:
        els.append(Text(0.6, H - 0.75, W - 1.2, 0.45, [caption], size=12, color=MUTED))
    return Slide(elements=els, notes=notes, title=title)


def stats_text(title, kicker, stats, items, notes="", size=17):
    els = header(title, kicker) + stats_row(stats)
    els.append(Text(0.6, 3.6, W - 1.2, H - 3.6 - 0.45, bullets(items, size), size=size, space_after=8))
    return Slide(elements=els, notes=notes, title=title)


def image_slide(title, kicker, img, caption=None, notes=""):
    els = header(title, kicker)
    top = 1.75 if kicker else 1.4
    bottom = H - (0.8 if caption else 0.4)
    els.append(Img(0.6, top, W - 1.2, bottom - top, str(img)))
    if caption:
        els.append(Text(0.6, H - 0.75, W - 1.2, 0.45, [caption], size=12, color=MUTED))
    return Slide(elements=els, notes=notes, title=title)


def table_slide(title, kicker, rows, col_w, items=None, highlight=None, notes="", size=12, row_h=0.36):
    els = header(title, kicker)
    top = 1.8
    els.append(Table(0.6, top, sum(col_w), rows, col_w, size=size, row_h=row_h, highlight_row=highlight))
    if items:
        ty = top + row_h * len(rows) + 0.3
        els.append(Text(0.6, ty, W - 1.2, H - ty - 0.4, bullets(items, 15), size=15, space_after=6))
    return Slide(elements=els, notes=notes, title=title)


def two_columns(title, kicker, left_title, left, right_title, right, notes="", size=15):
    els = header(title, kicker)
    top = 1.8
    cw = (W - 1.2 - 0.4) / 2
    for i, (t, items) in enumerate([(left_title, left), (right_title, right)]):
        x = 0.6 + i * (cw + 0.4)
        els.append(Box(x, top, cw, H - top - 0.45))
        els.append(Text(x + 0.2, top + 0.15, cw - 0.4, 0.55, [t], size=19, bold=True, color=INK))
        els.append(Text(x + 0.2, top + 0.75, cw - 0.4, H - top - 0.45 - 0.9, bullets(items, size), size=size, space_after=7))
    return Slide(elements=els, notes=notes, title=title)


def closing_slide(title, points, notes=""):
    els = [Text(0.8, 0.8, 11.7, 1.0, [title], size=34, color=WHITE, font="Cambria", bold=True, anchor="middle")]
    y = 2.1
    for i, p in enumerate(points, 1):
        els.append(Text(0.8, y, 0.8, 0.9, [str(i)], size=34, color=ICE, bold=True, anchor="middle"))
        els.append(Text(1.7, y, 10.8, 0.9, [p], size=18, color=WHITE, anchor="middle"))
        y += 1.2
    return Slide(elements=els, notes=notes, bg=DARK, title=title)
